# -*- coding: utf-8 -*-
"""Apply the final reference, citation and author changes to the Word build.

Input : 1019manuscript.docx  (the copy carrying the author title page)
Output: 1020manuscript.docx

Changes:
  * all seven figures replaced by the data-rebuilt PNGs in figures/,
    preserving their aspect ratios;
  * in-text [n] citations replaced by superscript numbers linked to bookmarks;
  * the numbered reference list rewritten in Nature style from verified DOI
    metadata, with a bookmark on every entry;
  * the author contributions paragraph populated for all nine authors.
"""

import copy
import io
import json
import os
import re
import sys

import docx
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml import parse_xml
from docx.shared import Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "1019manuscript.docx")
OUT = os.path.join(ROOT, "1020manuscript.docx")
REFS = os.path.join(ROOT, "work", "references_nc_1020.json")
FIGS = {n: os.path.join(ROOT, "figures", "Figure %d.png" % n)
        for n in range(1, 8)}
SVGS = {n: os.path.join(ROOT, "figures", "Figure %d.svg" % n)
        for n in range(1, 8)}

SVG_EXT_URI = "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"
SVG_NS = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"

CITE = re.compile(r"\[(\d+(?:\s*,\s*\d+)*)\]")
REF_HEAD = "References"
END_HEAD = "Data availability"

CONTRIB = (
    "Z.W. and L.W. contributed equally to this work and are co-first authors. "
    "Z.W. conceived and supervised the study, acquired funding and revised the "
    "manuscript. L.W. contributed to the study design, data interpretation and "
    "revised the manuscript. C.W. contributed to the formal analysis and revised "
    "the manuscript. X.X. contributed to the data curation and formal analysis. "
    "P.Z. contributed to the methodology and validation. Z.L. contributed to the "
    "investigation and visualization. H.F. contributed to the software and data "
    "curation. Q.H. contributed to the validation and revised the manuscript. "
    "M.G. contributed to the formal analysis and revised the manuscript. "
    "All authors reviewed and approved the manuscript."
)

TEXT_FIXES = [
    ("the shipped segmentation mask over sampled fields of view",
     "the shipped segmentation mask over all labelled fields of view"),
    ("certifies 11,741 cells", "certifies 11,742 cells"),
    ("A 20% budget commits 86.4% at 0.14 (Table 2).",
     "A 20% budget commits 86.4% at 0.14 (Fig. 5A)."),
    ("defers 44.1% of cells against 24.8% for the oracle, about 19 extra points",
     "defers 40.3% of cells against 24.4% for the oracle, about 16 extra points"),
    ("acceptance 0.85 at a 10% budget (Fig. 6A). Image-grouped splitting "
     "accepts 0.73 at selective error 0.08 but covers 0.66.",
     "acceptance 0.85 at a 10% budget. Image-grouped splitting "
     "accepts 0.73 at selective error 0.08 but covers 0.66 (Fig. 6A)."),
    ("for the within-intestine and transfer-to-tonsil-and-Barrett arms",
     "for the cross-FOV intestine and transfer-to-tonsil-and-Barrett arms"),
    ("Fraction of cells routed to review, and realised selective error among "
     "retained labels, versus nominal budget for the three split families.",
     "Fraction of cells routed to review versus nominal budget for the three "
     "split families; numbers give the realised selective error among retained "
     "labels at the 10% and 20% budgets."),
    ("Cross-field-of-view splitting already lowers coverage by about 8 points "
     "at a 10% budget.",
     "Cross-field-of-view splitting leaves coverage 0.82 at a 10% budget, "
     "8 points below nominal."),
]

TABLE_METHOD = {"crc": "certified", "crc_cond": "conditional",
                "naive": "uncalibrated", "oracle": "oracle"}


def clear_runs(p):
    for child in list(p._p):
        if child.tag != qn("w:pPr"):
            p._p.remove(child)


def set_run_font(run, size=11):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    rpr = run._r.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    rf.set(qn("w:eastAsia"), "Times New Roman")


def clone_run_text(src_run, text, superscript=False):
    r = copy.deepcopy(src_run._r)
    for child in list(r):
        if child.tag != qn("w:rPr"):
            r.remove(child)
    rpr = r.find(qn("w:rPr"))
    if superscript:
        if rpr is None:
            rpr = OxmlElement("w:rPr")
            r.insert(0, rpr)
        va = rpr.find(qn("w:vertAlign"))
        if va is None:
            va = OxmlElement("w:vertAlign")
            rpr.append(va)
        va.set(qn("w:val"), "superscript")
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    return r


def replace_citation_run(run):
    text = run.text
    matches = list(CITE.finditer(text))
    if not matches:
        return 0
    parent = run._r.getparent()
    idx = parent.index(run._r)
    elements = []
    last = 0
    for m in matches:
        if m.start() > last:
            elements.append(clone_run_text(run, text[last:m.start()]))
        nums = [int(x) for x in re.split(r"\s*,\s*", m.group(1))]
        for j, n in enumerate(nums):
            hl = OxmlElement("w:hyperlink")
            hl.set(qn("w:anchor"), "_Ref%d" % n)
            hl.set(qn("w:history"), "1")
            hl.append(clone_run_text(run, str(n), superscript=True))
            elements.append(hl)
            if j < len(nums) - 1:
                elements.append(clone_run_text(run, ",", superscript=True))
        last = m.end()
    if last < len(text):
        elements.append(clone_run_text(run, text[last:]))
    parent.remove(run._r)
    for off, el in enumerate(elements):
        parent.insert(idx + off, el)
    return len(matches)


def add_bookmark(p, name, bid):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bid))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bid))
    p._p.append(start)
    p._p.append(end)
    return start, end


def replace_reference_paragraph(p, n, text, bid):
    clear_runs(p)
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bid))
    start.set(qn("w:name"), "_Ref%d" % n)
    p._p.append(start)
    r = p.add_run("%d. %s" % (n, text))
    set_run_font(r, 11)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bid))
    p._p.append(end)
    p.paragraph_format.left_indent = Pt(24)
    p.paragraph_format.first_line_indent = Pt(-24)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(2)


def replace_figure(doc, caption_prefix, image_path):
    cap = next(i for i, p in enumerate(doc.paragraphs)
               if p.text.startswith(caption_prefix))
    prev = doc.paragraphs[cap - 1]
    if not prev._p.xpath(".//w:drawing"):
        raise AssertionError("the paragraph before %s is not an image"
                             % caption_prefix)
    idx = sum(1 for p in doc.paragraphs[:cap]
              if p._p.xpath(".//w:drawing")) - 1
    shape = doc.inline_shapes[idx]
    rid, image = doc.part.get_or_add_image(image_path)
    blip = shape._inline.graphic.graphicData.pic.blipFill.blip
    old_rid = blip.get(qn("r:embed"))
    blip.set(qn("r:embed"), rid)
    if old_rid and old_rid != rid:
        still_used = any(
            b.get(qn("r:embed")) == old_rid
            for b in doc.element.body.iter(qn("a:blip"))
        )
        if not still_used:
            doc.part.drop_rel(old_rid)
    w, h = image.px_width, image.px_height
    shape.width = Inches(6.2)
    shape.height = Inches(6.2 * h / w)
    attach_svg(doc, shape, os.path.splitext(image_path)[0] + ".svg")
    return idx, (w, h)


def attach_svg(doc, shape, svg_path):
    """Add an SVG part and point the picture's svgBlip at it.

    The existing PNG relationship stays in place as the required fallback for
    readers that do not understand DrawingML SVG extensions.
    """
    if not os.path.exists(svg_path):
        raise AssertionError("missing SVG for %s" % svg_path)
    blip = shape._inline.graphic.graphicData.pic.blipFill.blip
    png_rid = blip.get(qn("r:embed"))
    png_part = doc.part.related_parts[png_rid]
    svg_partname = PackURI(str(png_part.partname).rsplit(".", 1)[0] + ".svg")
    part = Part(svg_partname, "image/svg+xml",
                open(svg_path, "rb").read(), doc.part.package)
    svg_rid = doc.part.relate_to(part, RT.IMAGE)
    for old in blip.findall(qn("a:extLst")):
        blip.remove(old)
    blip.append(parse_xml(
        '<a:extLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
        ' xmlns:asvg="%s"'
        ' xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<a:ext uri="%s"><asvg:svgBlip r:embed="%s"/></a:ext></a:extLst>'
        % (SVG_NS, SVG_EXT_URI, svg_rid)))
    shape._inline.graphic.graphicData.pic.nvPicPr.cNvPr.set(
        "name", os.path.basename(svg_path))


def main():
    doc = docx.Document(SRC)
    refs = json.load(io.open(REFS, encoding="utf-8"))
    if len(refs) != 35:
        raise AssertionError("expected 35 references, got %d" % len(refs))

    doc_text = "\n".join(p.text for p in doc.paragraphs)
    old_share = "17.5% of cells, 1,794 of 10,283"
    new_share = "17.4% of cells, 1,794 of 10,283"
    fixed_share = 0
    if old_share in doc_text:
        for p in doc.paragraphs:
            for r in p.runs:
                if old_share in r.text:
                    r.text = r.text.replace(old_share, new_share)
                    fixed_share += 1
    elif new_share in doc_text:
        fixed_share = 1
    if fixed_share != 1:
        raise AssertionError("expected one delivery-share sentence, found %d"
                             % fixed_share)

    applied = {old: 0 for old, _ in TEXT_FIXES}
    for old, new in TEXT_FIXES:
        if old in doc_text:
            for p in doc.paragraphs:
                for r in p.runs:
                    if old in r.text:
                        applied[old] += r.text.count(old)
                        r.text = r.text.replace(old, new)
        elif new in doc_text:
            applied[old] = 1
    missed = [old for old, n in applied.items() if n != 1]
    if missed:
        raise AssertionError("text fixes not applied once: %r"
                             % [(m[:60], applied[m]) for m in missed])

    t1, t2 = doc.tables[0], doc.tables[1]
    for row in t1.rows[1:]:
        cell = row.cells[1]
        raw = cell.text.strip()
        if raw in TABLE_METHOD:
            for r in cell.paragraphs[0].runs:
                r.text = ""
            cell.paragraphs[0].runs[0].text = TABLE_METHOD[raw]
    for row in t2.rows:
        for cell in row.cells:
            if cell.text.strip() == "review (of pool)":
                cell.paragraphs[0].runs[0].text = "mean review across units"
            elif cell.text.strip() == "acceptance (before to after)":
                cell.paragraphs[0].runs[0].text = (
                    "mean acceptance (before to after)")
    note = next(p for p in doc.paragraphs
                if p.text.startswith("Notes: The certificate row"))
    old_note = "The certificate row is the source-only certificate of Table 1."
    new_note = (
        "The certificate row is the source-only certificate of Table 1. "
        "Acceptance and review are means of per-unit fractions, so each unit "
        "contributes equally. Fig. 5A uses pooled cell fractions; for the "
        "random split these are 65.3%, 40.3% and 13.6% at 5, 10 and 20% budgets.")
    hits = 1 if "Acceptance and review are means of per-unit fractions" \
        in note.text else 0
    if not hits:
        for r in note.runs:
            if old_note in r.text:
                r.text = r.text.replace(old_note, new_note)
                hits += 1
    if hits != 1:
        raise AssertionError("Table 2 note not updated")

    placed = {}
    for n in range(1, 8):
        placed[n] = replace_figure(doc, "Fig. %d |" % n, FIGS[n])

    fig2_caption = next(p for p in doc.paragraphs if p.text.startswith("Fig. 2 |"))
    old_h = (" Exchangeable and image-grouped splits leave calibration error "
             "unchanged; cross-tissue transfer multiplies it by 2.4 and costs "
             "19 points of accuracy.")
    new_h = (" Ratio of deployment-domain to training-domain ECE for each split "
             "family. Exchangeable and image-grouped splits leave calibration "
             "error relatively unchanged, whereas cross-tissue transfer "
             "multiplies it by about 2.4.")
    if old_h in fig2_caption.text:
        for run in fig2_caption.runs:
            if old_h in run.text:
                run.text = run.text.replace(old_h, new_h)
                break

    heads = [(i, p.text.strip()) for i, p in enumerate(doc.paragraphs)]
    ri = next(i for i, t in heads if t == REF_HEAD)
    ei = next(i for i, t in heads if t == END_HEAD)
    ref_paras = [p for p in doc.paragraphs[ri + 1:ei] if p.text.strip()]
    if len(ref_paras) != 35:
        raise AssertionError("reference section has %d entries" % len(ref_paras))

    old_n = [int(re.match(r"\[(\d+)\]", p.text).group(1)) for p in ref_paras]
    if old_n != list(range(1, 36)):
        raise AssertionError("reference paragraphs are not in 1..35 order")

    for p, ref in zip(ref_paras, refs):
        replace_reference_paragraph(p, ref["n"], ref["text"], 9000 + ref["n"])

    changes = 0
    for i, p in enumerate(doc.paragraphs):
        if ri <= i <= ei:
            continue
        for run in list(p.runs):
            changes += replace_citation_run(run)
    if changes < 10:
        raise AssertionError("too few citation groups replaced: %d" % changes)

    contrib = next(p for p in doc.paragraphs if p.text.startswith("<A.B.>"))
    clear_runs(contrib)
    r = contrib.add_run(CONTRIB)
    set_run_font(r, 11)

    doc.save(OUT)
    print("wrote %s" % OUT)
    print("  citation groups replaced:", changes)
    print("  references rewritten:", len(refs))
    for n in range(1, 8):
        idx, size = placed[n]
        print("  Figure %d: drawing %d, %dx%d px" % (n, idx + 1, size[0], size[1]))


if __name__ == "__main__":
    main()
