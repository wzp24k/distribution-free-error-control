# -*- coding: utf-8 -*-
"""Verify the rebuilt figures, their SVG text colour and their Word embeds."""

import glob
import hashlib
import io
import os
import re
import zipfile

import docx

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIGDIR = os.path.join(ROOT, "figures")
DOCX = os.path.join(ROOT, "1020manuscript.docx")


def main():
    svg_counts = {}
    for n in range(1, 8):
        path = os.path.join(FIGDIR, "Figure %d.svg" % n)
        xml = io.open(path, encoding="utf-8").read()
        nodes = re.findall(r"<text[^>]*>", xml)
        fills = []
        for node in nodes:
            m = re.search(r'fill="([^"]+)"', node)
            fills.append((m.group(1) if m else "inherit").lower())
        bad = [f for f in fills if f not in ("#000000", "black")]
        if bad:
            raise AssertionError("Figure %d has non-black text fills: %r"
                                 % (n, sorted(set(bad))))
        if not nodes:
            raise AssertionError("Figure %d has no editable text" % n)
        svg_counts[n] = len(nodes)

    doc = docx.Document(DOCX)
    if len(doc.inline_shapes) != 7:
        raise AssertionError("expected 7 inline figures, got %d"
                             % len(doc.inline_shapes))
    if len(doc.tables) != 3:
        raise AssertionError("expected 3 tables, got %d" % len(doc.tables))

    z = zipfile.ZipFile(DOCX)
    media = {hashlib.md5(z.read(n)).hexdigest(): n
             for n in z.namelist() if n.startswith("word/media/")}
    for n in range(1, 8):
        png = os.path.join(FIGDIR, "Figure %d.png" % n)
        h = hashlib.md5(open(png, "rb").read()).hexdigest()
        if h not in media:
            raise AssertionError("Figure %d PNG is not embedded in %s"
                                 % (n, os.path.basename(DOCX)))

    captions = [p.text for p in doc.paragraphs if p.text.startswith("Fig. ")]
    if len(captions) != 7:
        raise AssertionError("expected 7 figure captions, got %d"
                             % len(captions))
    if not all(("Fig. %d |" % n) in captions[n - 1] for n in range(1, 8)):
        raise AssertionError("figure captions are not in order")

    print("OK 7 figures, 3 tables")
    for n in range(1, 8):
        print("Figure %d: %d editable text elements, all black, PNG embedded"
              % (n, svg_counts[n]))


if __name__ == "__main__":
    main()
